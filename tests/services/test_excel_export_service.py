import io
from datetime import datetime
from decimal import Decimal

from bson import Decimal128, ObjectId
from openpyxl import load_workbook

from app.services.excel_export_service import (
    _converter_decimal,
    _formatar_data_texto,
    _formatar_valor_texto,
    gerar_excel_timeline,
)


def cco_mock():
    return {
        "_id": "CCO_TEST_001",
        "contratoCpp": "CONTRATO-001",
        "campo": "CAMPO_TESTE",
        "remessa": "REM-001",
        "faseRemessa": "FASE-01",
        "exercicio": 2024,
        "periodo": "2024-03",
        "anoReconhecimento": 2024,
        "mesReconhecimento": 3,
        "valorReconhecido": Decimal128("1234567.89"),
        "valorReconhecidoComOH": Decimal128("1300000.00"),
        "ipcaAcumuladoReais": Decimal128("10000.00"),
        "igpmAcumuladoReais": Decimal128("15000.00"),
        "diferencaValor": Decimal128("2000.00"),
        "overHeadTotal": Decimal128("10000.00"),
        "correcoesMonetarias": [
            {
                "tipo": "IPCA",
                "dataCorrecao": "2024-03-15T14:30:00Z",
                "taxaCorrecao": Decimal128("0.45"),
                "igpmAcumuladoReais": Decimal128("1000.50"),
                "valorReconhecido": Decimal128("1234567.89"),
                "valorReconhecidoComOH": Decimal128("1300000.00"),
                "overHeadTotal": Decimal128("10000.00"),
                "diferencaValor": Decimal128("1000.50"),
            },
            {
                "tipo": "IGPM",
                "dataCorrecao": "2024-04-15T10:00:00Z",
                "taxaCorrecao": Decimal128("0.40"),
                "igpmAcumuladoReais": Decimal128("900.10"),
                "valorReconhecido": Decimal128("1240000.00"),
                "valorReconhecidoComOH": Decimal128("1310000.00"),
                "overHeadTotal": Decimal128("10500.00"),
                "diferencaValor": Decimal128("900.10"),
            },
        ],
        "campoExtraString": "valor extra",
        "campoExtraObjectId": ObjectId("507f1f77bcf86cd799439011"),
        "campoExtraDict": {"chave": "valor"},
        "campoExtraLista": [1, 2, 3],
    }


def timeline_mock():
    return {
        "estado_temporal": False,
        "cco_id": "CCO_TEST_001",
        "data_corte": None,
        "eventos": [
            {
                "data": datetime(2024, 3, 15, 14, 30),
                "tipo": "CRIACAO",
                "titulo": "Criação da CCO",
                "descricao": "CCO criada para teste",
                "valor": Decimal128("1000.00"),
                "valor_acumulado": Decimal128("1000.00"),
                "valores": {
                    "valorReconhecido": Decimal128("1234567.89"),
                    "valorReconhecidoComOH": Decimal128("1300000.00"),
                    "diferencaValor": Decimal128("2000.00"),
                },
            },
            {
                "dataCorrecaoFormatada": "20/03/2024 09:00",
                "tipo": "IPCA",
                "titulo": "Correção IPCA",
                "descricao": "Correção monetária aplicada",
                "valores": {
                    "valorReconhecido": Decimal128("1240000.00"),
                    "valorReconhecidoComOH": Decimal128("1310000.00"),
                    "diferencaValor": Decimal128("3000.00"),
                },
            },
        ],
    }


def salvar_e_reabrir_workbook(workbook):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return load_workbook(output, data_only=False)


def test_gerar_excel_timeline_cria_quatro_abas():
    workbook = gerar_excel_timeline(cco_mock(), timeline_mock())

    assert workbook.sheetnames == [
        "Resumo",
        "Timeline",
        "Correções Monetárias",
        "Metadados",
    ]


def test_aba_resumo_preenche_dados_basicos():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Resumo"]

    assert ws["A1"].value == "TIMELINE DA CCO"
    assert ws["A4"].value == "ID da CCO:"
    assert ws["B4"].value == "CCO_TEST_001"
    assert ws["A5"].value == "Contrato CPP:"
    assert ws["B5"].value == "CONTRATO-001"


def test_aba_resumo_exibe_data_corte_e_estado_temporal():
    timeline = timeline_mock()
    timeline["estado_temporal"] = True
    timeline["data_corte"] = "2024-12-31"

    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline, data_corte="2024-12-31")
    )

    ws = workbook["Resumo"]

    assert ws["A2"].value == "Estado em: 2024-12-31"

    valores_resumo = {
        ws[f"A{row}"].value: ws[f"B{row}"].value
        for row in range(1, ws.max_row + 1)
    }

    assert valores_resumo["TOTAL DE EVENTOS"] == 2
    assert valores_resumo["Estado Temporal:"] == "Sim"


def test_aba_resumo_formata_valores_monetarios():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Resumo"]

    valores_resumo = {
        ws[f"A{row}"].value: (ws[f"B{row}"].value, ws[f"B{row}"].number_format)
        for row in range(1, ws.max_row + 1)
    }

    valor, formato = valores_resumo["Valor Reconhecido + OH:"]

    assert valor == 1300000.00
    assert formato == "R$ #,##0.00"


def test_aba_timeline_preenche_eventos():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Timeline"]

    assert ws["A1"].value == "Data"
    assert ws["B1"].value == "Tipo"
    assert ws["C1"].value == "Título"

    assert ws["B2"].value == "CRIACAO"
    assert ws["C2"].value == "Criação da CCO"
    assert ws["D2"].value == "CCO criada para teste"

    assert ws["B3"].value == "IPCA"
    assert ws["C3"].value == "Correção IPCA"


def test_aba_timeline_aplica_formatacao_monetaria():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Timeline"]

    assert ws["E2"].value == 1000.00
    assert ws["E2"].number_format == "R$ #,##0.00"
    assert ws["G2"].value == 1234567.89
    assert ws["G2"].number_format == "R$ #,##0.00"


def test_aba_timeline_aceita_lista_direta_de_eventos():
    eventos = timeline_mock()["eventos"]

    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), eventos)
    )

    ws = workbook["Timeline"]

    assert ws["B2"].value == "CRIACAO"
    assert ws["B3"].value == "IPCA"


def test_aba_correcoes_preenche_dados():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["A1"].value == "Tipo"
    assert ws["B1"].value == "Ano"
    assert ws["C1"].value == "Mês"
    assert ws["D1"].value == "Valor Reconhecido"
    assert ws["E1"].value == "Valor Rec. c/ Overhead"
    assert ws["F1"].value == "Overhead Total"
    assert ws["G1"].value == "Diferença"
    assert ws["H1"].value == "Taxa de Correção"
    assert ws["I1"].value == "IGPM Acumulado (R$)"
    assert ws["J1"].value == "Data Aplicação"

    assert ws["A2"].value == "IPCA"
    assert ws["B2"].value == 2024
    assert ws["C2"].value == 3
    assert ws["D2"].value == 1234567.89
    assert ws["E2"].value == 1300000.00
    assert ws["F2"].value == 10000.00
    assert ws["G2"].value == 1000.50
    assert ws["H2"].value == 0.45
    assert ws["I2"].value == 1000.50
    assert ws["J2"].value == "15/03/2024 14:30"

    assert ws["A3"].value == "IGPM"
    assert ws["B3"].value == 2024
    assert ws["C3"].value == 4
    assert ws["D3"].value == 1240000.00
    assert ws["E3"].value == 1310000.00
    assert ws["F3"].value == 10500.00
    assert ws["G3"].value == 900.10
    assert ws["H3"].value == 0.40
    assert ws["I3"].value == 900.10


def test_aba_correcoes_formatacao_percentual_e_monetaria():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["D2"].number_format == "R$ #,##0.00"
    assert ws["E2"].number_format == "R$ #,##0.00"
    assert ws["F2"].number_format == "R$ #,##0.00"
    assert ws["G2"].number_format == "R$ #,##0.00"
    assert ws["H2"].number_format == "0.00%"
    assert ws["I2"].number_format == "R$ #,##0.00"

    assert ws["D3"].number_format == "R$ #,##0.00"
    assert ws["E3"].number_format == "R$ #,##0.00"
    assert ws["F3"].number_format == "R$ #,##0.00"
    assert ws["G3"].number_format == "R$ #,##0.00"
    assert ws["H3"].number_format == "0.00%"
    assert ws["I3"].number_format == "R$ #,##0.00"


def test_aba_correcoes_taxa_exibida_como_percentual_do_fator():
    cco = cco_mock()
    cco["correcoesMonetarias"] = [
        {
            "tipo": "IPCA",
            "dataCorrecao": "2024-03-15T14:30:00Z",
            "taxaCorrecao": Decimal128("1.0590000"),
            "igpmAcumuladoReais": Decimal128("1000.50"),
            "valorReconhecido": Decimal128("1234567.89"),
            "valorReconhecidoComOH": Decimal128("1300000.00"),
            "overHeadTotal": Decimal128("10000.00"),
            "diferencaValor": Decimal128("1000.50"),
        }
    ]

    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco, timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["H2"].value == 1.059
    assert ws["H2"].number_format == "0.00%"


def test_aba_correcoes_sem_correcoes_nao_cria_linhas():
    cco = cco_mock()
    cco["correcoesMonetarias"] = []

    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco, timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws.max_row == 1
    assert ws["A1"].value == "Tipo"


def test_aba_correcoes_exibe_data_aplicacao_formatada():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["J2"].value == "15/03/2024 14:30"
    assert ws["J3"].value == "15/04/2024 10:00"


def test_aba_correcoes_exibe_tipo_correto():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["A2"].value == "IPCA"
    assert ws["A3"].value == "IGPM"


def test_aba_correcoes_tipo_retificacao_exibir_dados():
    cco = cco_mock()
    cco["correcoesMonetarias"] = [
        {
            "tipo": "RETIFICACAO",
            "dataCorrecao": "2024-05-10T08:00:00Z",
            "taxaCorrecao": Decimal128("0.10"),
            "igpmAcumuladoReais": Decimal128("300.00"),
            "valorReconhecido": Decimal128("100.00"),
            "valorReconhecidoComOH": Decimal128("110.00"),
            "overHeadTotal": Decimal128("10.00"),
            "diferencaValor": Decimal128("20.00"),
        }
    ]

    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco, timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["A2"].value == "RETIFICACAO"
    assert ws["B2"].value == 2024
    assert ws["C2"].value == 5
    assert ws["D2"].value == 100.00
    assert ws["H2"].value == 0.10
    assert ws["I2"].value == 300.00


def test_aba_correcoes_tipo_recuperacao_exibir_dados():
    cco = cco_mock()
    cco["correcoesMonetarias"] = [
        {
            "tipo": "RECUPERACAO",
            "dataCorrecao": "2024-06-20T09:00:00Z",
            "taxaCorrecao": Decimal128("1.0200000"),
            "igpmAcumuladoReais": Decimal128("450.00"),
            "valorReconhecido": Decimal128("5000.00"),
            "valorReconhecidoComOH": Decimal128("5500.00"),
            "overHeadTotal": Decimal128("500.00"),
            "diferencaValor": Decimal128("700.00"),
        }
    ]

    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco, timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["A2"].value == "RECUPERACAO"
    assert ws["B2"].value == 2024
    assert ws["C2"].value == 6
    assert ws["D2"].value == 5000.00
    assert ws["E2"].value == 5500.00
    assert ws["F2"].value == 500.00
    assert ws["G2"].value == 700.00
    assert ws["H2"].value == 1.02
    assert ws["I2"].value == 450.00


def test_aba_correcoes_data_correcao_none_gera_campos_vazios():
    cco = cco_mock()
    cco["correcoesMonetarias"] = [
        {
            "tipo": "IPCA",
            "dataCorrecao": None,
            "taxaCorrecao": Decimal128("0.30"),
            "igpmAcumuladoReais": Decimal128("0"),
            "valorReconhecido": Decimal128("100.00"),
            "valorReconhecidoComOH": Decimal128("110.00"),
            "overHeadTotal": Decimal128("10.00"),
            "diferencaValor": Decimal128("20.00"),
        }
    ]

    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco, timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["B2"].value is None
    assert ws["C2"].value is None
    assert ws["J2"].value is None
    assert ws["A2"].value == "IPCA"


def test_aba_correcoes_data_correcao_invalida_gera_campos_vazios():
    cco = cco_mock()
    cco["correcoesMonetarias"] = [
        {
            "tipo": "IGPM",
            "dataCorrecao": "data-invalida",
            "taxaCorrecao": Decimal128("0.25"),
            "igpmAcumuladoReais": Decimal128("150.00"),
            "valorReconhecido": Decimal128("100.00"),
            "valorReconhecidoComOH": Decimal128("110.00"),
            "overHeadTotal": Decimal128("10.00"),
            "diferencaValor": Decimal128("20.00"),
        }
    ]

    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco, timeline_mock())
    )

    ws = workbook["Correções Monetárias"]

    assert ws["B2"].value is None
    assert ws["C2"].value is None
    assert ws["I2"].value == 150.00
    assert ws["A2"].value == "IGPM"


def test_aba_metadados_lista_campos_e_ignora_correcoes():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Metadados"]

    valores_coluna_a = [
        ws[f"A{row}"].value
        for row in range(1, ws.max_row + 1)
    ]

    assert "METADADOS COMPLETOS" in valores_coluna_a
    assert "contratoCpp" in valores_coluna_a
    assert "campoExtraString" in valores_coluna_a
    assert "correcoesMonetarias" not in valores_coluna_a


def test_aba_metadados_serializa_tipos_complexos():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Metadados"]

    metadados = {
        ws[f"A{row}"].value: ws[f"B{row}"].value
        for row in range(1, ws.max_row + 1)
    }

    assert metadados["campoExtraObjectId"] == "507f1f77bcf86cd799439011"
    assert '"chave": "valor"' in metadados["campoExtraDict"]
    assert metadados["campoExtraLista"] == "[1, 2, 3]"


def test_workbook_pode_ser_salvo_em_bytesio():
    workbook = gerar_excel_timeline(cco_mock(), timeline_mock())

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    assert output.getbuffer().nbytes > 0


def test_converter_decimal_trata_decimal128_decimal_numero_none_e_invalido():
    assert _converter_decimal(Decimal128("10.50")) == 10.50
    assert _converter_decimal(Decimal("20.75")) == 20.75
    assert _converter_decimal("30.25") == 30.25
    assert _converter_decimal(None) == 0
    assert _converter_decimal("abc") == 0


def test_formatar_data_texto_trata_datetime_date_string_iso_none_e_string_invalida():
    assert _formatar_data_texto(datetime(2024, 3, 15, 14, 30)) == "15/03/2024 14:30"
    assert _formatar_data_texto(datetime(2024, 3, 15).date()) == "15/03/2024"
    assert _formatar_data_texto("2024-03-15T14:30:00") == "15/03/2024 14:30"
    assert _formatar_data_texto(None) == ""
    assert _formatar_data_texto("data-invalida") == "data-invalida"


def test_formatar_valor_texto_trata_tipos_complexos():
    object_id = ObjectId("507f1f77bcf86cd799439011")

    assert _formatar_valor_texto(Decimal128("10.50")) == 10.50
    assert _formatar_valor_texto(Decimal("20.75")) == 20.75
    assert _formatar_valor_texto(object_id) == "507f1f77bcf86cd799439011"
    assert '"a": 1' in _formatar_valor_texto({"a": 1})
    assert _formatar_valor_texto([1, 2]) == "[1, 2]"
    assert _formatar_valor_texto(None) == ""
