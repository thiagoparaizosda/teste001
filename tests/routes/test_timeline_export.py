import io
from datetime import datetime
from unittest.mock import MagicMock

import pytest
from bson import Decimal128
from openpyxl import load_workbook, workbook


CCO_ID = "CCO_TEST_001"


@pytest.fixture
def cco_mock():
    return {
        "_id": CCO_ID,
        "contratoCpp": "CONTRATO-001",
        "campo": "CAMPO_TESTE",
        "remessa": "REM-001",
        "faseRemessa": "FASE-01",
        "periodo": "2024-03",
        "anoReconhecimento": 2024,
        "mesReconhecimento": 3,
    }


@pytest.fixture
def timeline_mock():
    return {
        "estado_temporal": False,
        "cco_id": CCO_ID,
        "data_corte": None,
        "eventos": [
            {
                "tipo": "CRIACAO",
                "titulo": "Criação da CCO",
                "descricao": "CCO criada para teste",
                "dataCorrecaoFormatada": "15/03/2024 14:30",
                "valores": {
                    "valorReconhecido": Decimal128("1234567.89"),
                    "valorReconhecidoComOH": Decimal128("1300000.00"),
                    "overHeadTotal": Decimal128("10000.00"),
                    "diferencaValor": Decimal128("2000.00"),
                    "valorRecuperado": Decimal128("300.00"),
                    "taxaCorrecao": 0.1234,
                },
                "detalhes": {"ativo": True},
            }
        ],
        "valores_em_data_corte": {},
    }


def configurar_db_mock(monkeypatch, cco_mock):
    from app.routes import portal_ui

    collection_mock = MagicMock()
    collection_mock.find_one.return_value = cco_mock

    db_mock = MagicMock()
    db_mock.conta_custo_oleo_entity = collection_mock

    monkeypatch.setattr(portal_ui.portal_service, "_get_db_prd", lambda: db_mock)

    return collection_mock


def test_export_timeline_xlsx_sucesso(client, monkeypatch, cco_mock, timeline_mock):
    from app.routes import portal_ui

    configurar_db_mock(monkeypatch, cco_mock)

    monkeypatch.setattr(
        portal_ui.portal_service,
        "get_timeline_with_cutoff",
        lambda cco_id, data_corte=None: timeline_mock,
        raising=False,
    )

    monkeypatch.setattr(
        portal_ui,
        "verificar_rate_limit_export",
        lambda chave, limite=10, janela_segundos=60: True,
    )

    response = client.get(f"/api/cco-timeline/{CCO_ID}/export")

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in response.headers["Content-Disposition"]
    assert ".xlsx" in response.headers["Content-Disposition"]

    workbook = load_workbook(io.BytesIO(response.data))

    assert "Resumo" in workbook.sheetnames
    assert "Timeline" in workbook.sheetnames
    assert "Correções Monetárias" in workbook.sheetnames
    assert "Metadados" in workbook.sheetnames


def test_export_timeline_formato_nao_suportado(client):
    response = client.get(f"/api/cco-timeline/{CCO_ID}/export?format=csv")

    assert response.status_code == 400
    assert response.get_json()["error"] == "Formato não suportado"


def test_export_timeline_data_corte_invalida(client, monkeypatch):
    from app.routes import portal_ui

    monkeypatch.setattr(
        portal_ui,
        "verificar_rate_limit_export",
        lambda chave, limite=10, janela_segundos=60: True,
    )

    response = client.get(f"/api/cco-timeline/{CCO_ID}/export?data_corte=31/12/2024")

    assert response.status_code == 400
    assert response.get_json()["error"] == "Formato de data inválido"


def test_export_timeline_cco_nao_encontrada(client, monkeypatch):
    from app.routes import portal_ui

    configurar_db_mock(monkeypatch, None)

    monkeypatch.setattr(
        portal_ui,
        "verificar_rate_limit_export",
        lambda chave, limite=10, janela_segundos=60: True,
    )

    response = client.get(f"/api/cco-timeline/{CCO_ID}/export")

    assert response.status_code == 404
    assert response.get_json()["error"] == "CCO não encontrada"


def test_export_timeline_erro_no_service(client, monkeypatch, cco_mock):
    from app.routes import portal_ui

    configurar_db_mock(monkeypatch, cco_mock)

    monkeypatch.setattr(
        portal_ui.portal_service,
        "get_timeline_with_cutoff",
        lambda cco_id, data_corte=None: {"erro": "Erro ao buscar timeline"},
        raising=False,
    )

    monkeypatch.setattr(
        portal_ui,
        "verificar_rate_limit_export",
        lambda chave, limite=10, janela_segundos=60: True,
    )

    response = client.get(f"/api/cco-timeline/{CCO_ID}/export")

    assert response.status_code == 404
    assert response.get_json()["erro"] == "Erro ao buscar timeline"


def test_export_timeline_limite_eventos(client, monkeypatch, cco_mock):
    from app.routes import portal_ui

    configurar_db_mock(monkeypatch, cco_mock)

    timeline_grande = {
        "estado_temporal": False,
        "cco_id": CCO_ID,
        "data_corte": None,
        "eventos": [{"tipo": "CRIACAO", "valores": {}, "detalhes": {}} for _ in range(1001)],
        "valores_em_data_corte": {},
    }

    monkeypatch.setattr(
        portal_ui.portal_service,
        "get_timeline_with_cutoff",
        lambda cco_id, data_corte=None: timeline_grande,
        raising=False,
    )

    monkeypatch.setattr(
        portal_ui,
        "verificar_rate_limit_export",
        lambda chave, limite=10, janela_segundos=60: True,
        raising=False,
    )

    response = client.get(f"/api/cco-timeline/{CCO_ID}/export")

    assert response.status_code == 400
    assert response.get_json()["error"] == "Timeline muito grande para exportação"


def test_export_timeline_rate_limit(client, monkeypatch):
    from app.routes import portal_ui

    monkeypatch.setattr(
        portal_ui,
        "verificar_rate_limit_export",
        lambda chave, limite=10, janela_segundos=60: False,
    )

    response = client.get(f"/api/cco-timeline/{CCO_ID}/export")

    assert response.status_code == 429
    assert response.get_json()["error"] == "Limite de exportações excedido"


def test_export_timeline_erro_interno(client, monkeypatch, cco_mock):
    from app.routes import portal_ui

    configurar_db_mock(monkeypatch, cco_mock)

    def gerar_erro(cco_id, data_corte=None):
        raise Exception("Erro simulado")

    monkeypatch.setattr(
        portal_ui.portal_service,
        "get_timeline_with_cutoff",
        gerar_erro,
        raising=False,
    )

    monkeypatch.setattr(
        portal_ui,
        "verificar_rate_limit_export",
        lambda chave, limite=10, janela_segundos=60: True,
    )

    response = client.get(f"/api/cco-timeline/{CCO_ID}/export")

    assert response.status_code == 500
    assert response.get_json()["error"] == "Erro interno ao exportar timeline"

def test_export_timeline_nome_arquivo_sem_id_e_com_campos(client, monkeypatch, cco_mock, timeline_mock):
    from app.routes import portal_ui

    cco_mock.update({
        "campo": "ARAM",
        "remessa": 30,
        "faseRemessa": "MEN",
        "anoReconhecimento": 2023,
        "mesReconhecimento": 3,
        "periodo": 1,
        "origemDosGastos": "GASTO_EXCLUSIVO",
    })

    configurar_db_mock(monkeypatch, cco_mock)

    monkeypatch.setattr(
        portal_ui.portal_service,
        "get_timeline_with_cutoff",
        lambda cco_id, data_corte=None: timeline_mock,
        raising=False,
    )

    monkeypatch.setattr(
        portal_ui,
        "verificar_rate_limit_export",
        lambda chave, limite=10, janela_segundos=60: True,
    )

    response = client.get(f"/api/cco-timeline/{CCO_ID}/export")

    assert response.status_code == 200

    content_disposition = response.headers["Content-Disposition"]

    assert ".xlsx" in content_disposition
    assert CCO_ID not in content_disposition

    assert "ARAM" in content_disposition
    assert "30" in content_disposition
    assert "MEN" in content_disposition
    assert "2023-03" in content_disposition
    assert "_1" in content_disposition
