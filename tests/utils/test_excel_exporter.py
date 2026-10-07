import io
from datetime import datetime
from bson import Decimal128
from openpyxl import load_workbook

from app.utils.excel_exporter import gerar_excel_timeline, formatar_valor_excel


def cco_mock():
    return {
        "_id": "CCO_TEST_001",
        "contratoCpp": "CONTRATO-001",
        "campo": "CAMPO_TESTE",
        "remessa": "REM-001",
        "faseRemessa": "FASE-01",
        "periodo": "2024-03",
        "anoReconhecimento": 2024,
        "mesReconhecimento": 3,
    }


def timeline_mock(data_corte=None):
    return {
        "estado_temporal": bool(data_corte),
        "cco_id": "CCO_TEST_001",
        "data_corte": data_corte,
        "eventos": [
            {
                "tipo": "CRIACAO",
                "titulo": "Criação da CCO",
                "descricao": "CCO criada para teste",
                "dataCorrecaoFormatada": "15/03/2024 14:30",
                "valores": {
                    "valorReconhecido": Decimal128("1234567.89"),
                    "valorReconhecidoComOH": Decimal128("1300000.00"),
                    "overHeadTotal": Decimal128("10000.00"),
                    "diferencaValor": Decimal128("2000.00"),
                    "valorRecuperado": Decimal128("300.00"),
                    "taxaCorrecao": 0.1234,
                },
                "detalhes": {
                    "ativo": True,
                    "observacao": "Evento de teste",
                },
            }
        ],
        "valores_em_data_corte": {
            "valorReconhecido": Decimal128("1234567.89"),
            "valorReconhecidoComOH": Decimal128("1300000.00"),
            "overHeadTotal": Decimal128("10000.00"),
            "diferencaValor": Decimal128("2000.00"),
        },
    }


def salvar_e_reabrir_workbook(workbook):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return load_workbook(output)


def test_gerar_excel_timeline_cria_abas():
    workbook = gerar_excel_timeline(cco_mock(), timeline_mock())

    assert "Resumo" in workbook.sheetnames
    assert "Timeline" in workbook.sheetnames
    assert "Valores" in workbook.sheetnames


def test_gerar_excel_timeline_preenche_resumo():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Resumo"]

    assert ws["A1"].value == "Campo"
    assert ws["B1"].value == "Valor"
    assert ws["A2"].value == "ID da CCO"
    assert ws["B2"].value == "CCO_TEST_001"


def test_gerar_excel_timeline_preenche_eventos():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock())
    )

    ws = workbook["Timeline"]

    assert ws["A1"].value == "Ordem"
    assert ws["B1"].value == "Tipo"
    assert ws["B2"].value == "CRIACAO"
    assert ws["C2"].value == "Criação da CCO"


def test_gerar_excel_timeline_com_data_corte():
    workbook = salvar_e_reabrir_workbook(
        gerar_excel_timeline(cco_mock(), timeline_mock("2024-12-31"), data_corte="2024-12-31")
    )

    ws_resumo = workbook["Resumo"]

    valores_resumo = {
        ws_resumo[f"A{row}"].value: ws_resumo[f"B{row}"].value
        for row in range(1, ws_resumo.max_row + 1)
    }

    assert valores_resumo["Data de corte"] == "2024-12-31"
    assert valores_resumo["Estado temporal"] is True


def test_gerar_excel_timeline_salva_em_bytesio():
    workbook = gerar_excel_timeline(cco_mock(), timeline_mock())

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    assert output.getbuffer().nbytes > 0


def test_formatar_valor_excel_decimal128():
    assert formatar_valor_excel(Decimal128("1234567.89")) == 1234567.89


def test_formatar_valor_excel_datetime():
    data = datetime(2024, 3, 15, 14, 30)

    assert formatar_valor_excel(data) == "15/03/2024 14:30"


def test_formatar_valor_excel_dict_lista():
    valor_dict = {"chave": "valor"}
    valor_lista = [1, 2, 3]

    assert '"chave": "valor"' in formatar_valor_excel(valor_dict)
    assert "[1, 2, 3]" == formatar_valor_excel(valor_lista)


def test_formatar_valor_excel_none():
    assert formatar_valor_excel(None) == ""
