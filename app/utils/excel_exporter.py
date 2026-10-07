import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


def formatar_valor_excel(valor):
    """
    Converte valores para uma representação segura no Excel.
    Trata Decimal128, datetime, listas, dicts e None.
    """
    if valor is None:
        return ""

    if hasattr(valor, "to_decimal"):
        return float(valor.to_decimal())

    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")

    if isinstance(valor, (dict, list)):
        return json.dumps(valor, ensure_ascii=False, default=str)

    return valor


def ajustar_largura_colunas(ws):
    for coluna in ws.columns:
        max_length = 0
        coluna_letra = get_column_letter(coluna[0].column)

        for cell in coluna:
            valor = cell.value
            if valor is not None:
                max_length = max(max_length, len(str(valor)))

        largura = min(max(max_length + 2, 12), 45)
        ws.column_dimensions[coluna_letra].width = largura


def estilizar_cabecalho(ws, linha=1):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="D9EAF7"),
        right=Side(style="thin", color="D9EAF7"),
        top=Side(style="thin", color="D9EAF7"),
        bottom=Side(style="thin", color="D9EAF7"),
    )

    for cell in ws[linha]:
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def aplicar_bordas(ws):
    border = Border(
        left=Side(style="thin", color="D9EAF7"),
        right=Side(style="thin", color="D9EAF7"),
        top=Side(style="thin", color="D9EAF7"),
        bottom=Side(style="thin", color="D9EAF7"),
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def gerar_excel_timeline(cco, timeline_resultado, data_corte=None):
    """
    Gera workbook Excel com dados da CCO e timeline.
    """
    wb = Workbook()

    ws_resumo = wb.active
    ws_resumo.title = "Resumo"

    ws_timeline = wb.create_sheet("Timeline")
    ws_valores = wb.create_sheet("Valores")

    # =========================
    # Aba Resumo
    # =========================
    ws_resumo.append(["Campo", "Valor"])
    estilizar_cabecalho(ws_resumo)

    resumo = [
        ("ID da CCO", cco.get("_id")),
        ("Contrato CPP", cco.get("contratoCpp")),
        ("Campo", cco.get("campo")),
        ("Remessa", cco.get("remessa")),
        ("Fase Remessa", cco.get("faseRemessa")),
        ("Período", cco.get("periodo")),
        ("Ano Reconhecimento", cco.get("anoReconhecimento")),
        ("Mês Reconhecimento", cco.get("mesReconhecimento")),
        ("Data de corte", data_corte if data_corte else "Não informada"),
        ("Estado temporal", timeline_resultado.get("estado_temporal")),
        ("Total de eventos", len(timeline_resultado.get("eventos", []))),
        ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
    ]

    for campo, valor in resumo:
        ws_resumo.append([campo, formatar_valor_excel(valor)])

    aplicar_bordas(ws_resumo)
    ajustar_largura_colunas(ws_resumo)

    # =========================
    # Aba Timeline
    # =========================
    headers_timeline = [
        "Ordem",
        "Tipo",
        "Título",
        "Descrição",
        "Data Correção",
        "Valor Reconhecido",
        "Valor Reconhecido c/ OH",
        "Overhead Total",
        "Diferença Valor",
        "Valor Recuperado",
        "Taxa Correção",
        "Detalhes",
    ]

    ws_timeline.append(headers_timeline)
    estilizar_cabecalho(ws_timeline)

    eventos = timeline_resultado.get("eventos", [])

    for index, evento in enumerate(eventos, start=1):
        valores = evento.get("valores", {}) or {}
        detalhes = evento.get("detalhes", {}) or {}

        ws_timeline.append([
            index,
            evento.get("tipo", ""),
            evento.get("titulo", ""),
            evento.get("descricao", ""),
            evento.get("dataCorrecaoFormatada") or formatar_valor_excel(evento.get("dataCorrecao")),
            formatar_valor_excel(valores.get("valorReconhecido")),
            formatar_valor_excel(valores.get("valorReconhecidoComOH")),
            formatar_valor_excel(valores.get("overHeadTotal")),
            formatar_valor_excel(valores.get("diferencaValor")),
            formatar_valor_excel(valores.get("valorRecuperado")),
            formatar_valor_excel(valores.get("taxaCorrecao")),
            formatar_valor_excel(detalhes),
        ])

    aplicar_bordas(ws_timeline)
    ajustar_largura_colunas(ws_timeline)

    for col in ["F", "G", "H", "I", "J"]:
        for cell in ws_timeline[col][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = 'R$ #,##0.00'

    for cell in ws_timeline["K"][1:]:
        if isinstance(cell.value, (int, float)):
            cell.number_format = '0.0000%'

    # =========================
    # Aba Valores
    # =========================
    ws_valores.append(["Campo", "Valor"])
    estilizar_cabecalho(ws_valores)

    valores_em_data = timeline_resultado.get("valores_em_data_corte", {}) or {}

    if valores_em_data:
        for campo, valor in valores_em_data.items():
            ws_valores.append([
                campo,
                formatar_valor_excel(valor)
            ])
    else:
        ws_valores.append(["Informação", "Nenhum valor de data de corte retornado"])

    aplicar_bordas(ws_valores)
    ajustar_largura_colunas(ws_valores)

    for row in ws_valores.iter_rows(min_row=2, min_col=2, max_col=2):
        cell = row[0]
        if isinstance(cell.value, (int, float)):
            cell.number_format = 'R$ #,##0.00'

    ws_resumo.freeze_panes = "A2"
    ws_timeline.freeze_panes = "A2"
    ws_valores.freeze_panes = "A2"

    return wb