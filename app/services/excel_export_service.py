import json
from datetime import date, datetime
from decimal import Decimal

from bson import Decimal128, ObjectId
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MOEDA_FORMAT = 'R$ #,##0.00'
PERCENTUAL_FORMAT = '0.00%'
DATA_FORMAT = "DD/MM/YYYY"
DATA_HORA_FORMAT = "DD/MM/YYYY HH:MM"


def _converter_decimal(valor):
    """
    Converte Decimal128, Decimal, int, float e strings numéricas para float.
    Retorna 0 quando o valor for nulo ou inválido.
    """
    if valor is None or valor == "":
        return 0

    if isinstance(valor, Decimal128):
        return float(valor.to_decimal())

    if isinstance(valor, Decimal):
        return float(valor)

    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0


def _formatar_data_texto(valor):
    """
    Formata datetime/date/string ISO para texto DD/MM/YYYY ou DD/MM/YYYY HH:MM.
    """
    if valor is None or valor == "":
        return ""

    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")

    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")

    if isinstance(valor, str):
        texto = valor.strip()

        if not texto:
            return ""

        try:
            data = datetime.fromisoformat(texto.replace("Z", "+00:00"))
            return data.strftime("%d/%m/%Y %H:%M")
        except ValueError:
            return texto

    return str(valor)


def _formatar_valor_texto(valor):
    """
    Converte valores complexos para representação segura no Excel.
    """
    if valor is None:
        return ""

    if isinstance(valor, Decimal128):
        return _converter_decimal(valor)

    if isinstance(valor, Decimal):
        return float(valor)

    if isinstance(valor, (datetime, date)):
        return _formatar_data_texto(valor)

    if isinstance(valor, ObjectId):
        return str(valor)

    if isinstance(valor, (dict, list)):
        return json.dumps(valor, ensure_ascii=False, default=str)

    return valor


def _aplicar_borda_e_alinhamento(ws):
    borda = Border(
        left=Side(style="thin", color="D9EAF7"),
        right=Side(style="thin", color="D9EAF7"),
        top=Side(style="thin", color="D9EAF7"),
        bottom=Side(style="thin", color="D9EAF7"),
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.border = borda
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _estilizar_cabecalho(ws, linha=1, cor="4472C4"):
    preenchimento = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
    fonte = Font(bold=True, color="FFFFFF")

    for cell in ws[linha]:
        cell.fill = preenchimento
        cell.font = fonte
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ajustar_largura_colunas(ws, largura_minima=12, largura_maxima=60):
    for coluna in ws.columns:
        max_length = 0
        coluna_letra = get_column_letter(coluna[0].column)

        for cell in coluna:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[coluna_letra].width = min(
            max(max_length + 2, largura_minima),
            largura_maxima
        )


def _obter_eventos(timeline):
    """
    Aceita timeline no formato {"eventos": [...]} ou lista direta.
    """
    if isinstance(timeline, dict):
        return timeline.get("eventos", []) or []

    if isinstance(timeline, list):
        return timeline

    return []


def _obter_valor_evento(evento, chave, padrao=0):
    """
    Busca valor tanto na raiz do evento quanto em evento["valores"].
    """
    if chave in evento:
        return evento.get(chave, padrao)

    valores = evento.get("valores", {}) or {}
    return valores.get(chave, padrao)


def gerar_excel_timeline(cco, timeline, data_corte=None):
    """
    Gera arquivo Excel com timeline da CCO.

    Estrutura:
        - Aba 1: Resumo
        - Aba 2: Timeline de Eventos
        - Aba 3: Correções Monetárias
        - Aba 4: Metadados Completos

    Args:
        cco (dict): documento da CCO.
        timeline (dict|list): timeline processada. Pode conter chave "eventos".
        data_corte (str|None): data de corte opcional no formato YYYY-MM-DD.

    Returns:
        Workbook: workbook openpyxl pronto para ser salvo.
    """
    wb = Workbook()

    aba_resumo = wb.active
    aba_resumo.title = "Resumo"

    aba_timeline = wb.create_sheet("Timeline")
    aba_correcoes = wb.create_sheet("Correções Monetárias")
    aba_metadados = wb.create_sheet("Metadados")

    _preencher_aba_resumo(aba_resumo, cco, timeline, data_corte)
    _preencher_aba_timeline(aba_timeline, timeline)
    _preencher_aba_correcoes(aba_correcoes, cco.get("correcoesMonetarias", []))
    _preencher_aba_metadados(aba_metadados, cco)

    return wb


def _preencher_aba_resumo(ws, cco, timeline, data_corte):
    """Aba 1: Resumo da CCO."""
    ws["A1"] = "TIMELINE DA CCO"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")

    if data_corte:
        ws["A2"] = f"Estado em: {data_corte}"
        ws["A2"].font = Font(italic=True, color="FF0000")

    linha = 4
    dados_basicos = [
        ("ID da CCO:", cco.get("_id")),
        ("Contrato CPP:", cco.get("contratoCpp")),
        ("Campo:", cco.get("campo")),
        ("Remessa:", cco.get("remessa")),
        ("Fase:", cco.get("faseRemessa")),
        ("Exercício:", cco.get("exercicio")),
        ("Período:", cco.get("periodo")),
        ("Ano Reconhecimento:", cco.get("anoReconhecimento")),
        ("Mês Reconhecimento:", cco.get("mesReconhecimento")),
    ]

    for label, valor in dados_basicos:
        ws[f"A{linha}"] = label
        ws[f"A{linha}"].font = Font(bold=True)
        ws[f"B{linha}"] = _formatar_valor_texto(valor)
        linha += 1

    linha = 15
    ws[f"A{linha}"] = "VALORES FINANCEIROS"
    ws[f"A{linha}"].font = Font(size=14, bold=True, color="1F4E78")
    linha += 1

    valores = [
        ("Valor Reconhecido:", cco.get("valorReconhecido"), MOEDA_FORMAT),
        ("Valor Reconhecido + OH:", cco.get("valorReconhecidoComOH"), MOEDA_FORMAT),
        ("IPCA Acumulado:", cco.get("ipcaAcumuladoReais", cco.get("ipca", 0)), MOEDA_FORMAT),
        ("IGPM Acumulado:", cco.get("igpmAcumuladoReais", cco.get("igpm", 0)), MOEDA_FORMAT),
        ("Diferença Total:", cco.get("diferencaValor", 0), MOEDA_FORMAT),
        ("Overhead Total:", cco.get("overHeadTotal", 0), MOEDA_FORMAT),
    ]

    for label, valor, formato in valores:
        ws[f"A{linha}"] = label
        ws[f"A{linha}"].font = Font(bold=True)

        cell_valor = ws[f"B{linha}"]
        cell_valor.value = _converter_decimal(valor)
        cell_valor.number_format = formato
        linha += 1

    eventos = _obter_eventos(timeline)

    linha += 1
    ws[f"A{linha}"] = "TOTAL DE EVENTOS"
    ws[f"A{linha}"].font = Font(bold=True)
    ws[f"B{linha}"] = len(eventos)

    if isinstance(timeline, dict) and timeline.get("estado_temporal") is not None:
        linha += 1
        ws[f"A{linha}"] = "Estado Temporal:"
        ws[f"A{linha}"].font = Font(bold=True)
        ws[f"B{linha}"] = "Sim" if timeline.get("estado_temporal") else "Não"

    _aplicar_borda_e_alinhamento(ws)
    _ajustar_largura_colunas(ws)
    ws.freeze_panes = "A4"


def _preencher_aba_timeline(ws, timeline):
    """Aba 2: Timeline de eventos cronológicos."""
    headers = [
        "Data",
        "Tipo",
        "Título",
        "Descrição",
        "Valor",
        "Valor Acumulado",
        "Valor Reconhecido",
        "Valor Reconhecido + OH",
        "Diferença",
    ]

    ws.append(headers)
    _estilizar_cabecalho(ws, linha=1, cor="4472C4")

    eventos = _obter_eventos(timeline)

    for row, evento in enumerate(eventos, start=2):
        data_evento = (
            evento.get("data")
            or evento.get("dataCorrecao")
            or evento.get("dataCorrecaoFormatada")
            or evento.get("dataCriacaoCorrecao")
        )

        valor = _obter_valor_evento(evento, "valor", 0)
        valor_acumulado = _obter_valor_evento(evento, "valor_acumulado", 0)
        valor_reconhecido = _obter_valor_evento(evento, "valorReconhecido", 0)
        valor_reconhecido_oh = _obter_valor_evento(evento, "valorReconhecidoComOH", 0)
        diferenca = _obter_valor_evento(evento, "diferencaValor", 0)

        ws.cell(row=row, column=1, value=_formatar_data_texto(data_evento))
        ws.cell(row=row, column=2, value=evento.get("tipo", ""))
        ws.cell(row=row, column=3, value=evento.get("titulo", ""))
        ws.cell(row=row, column=4, value=evento.get("descricao", ""))
        ws.cell(row=row, column=5, value=_converter_decimal(valor))
        ws.cell(row=row, column=6, value=_converter_decimal(valor_acumulado))
        ws.cell(row=row, column=7, value=_converter_decimal(valor_reconhecido))
        ws.cell(row=row, column=8, value=_converter_decimal(valor_reconhecido_oh))
        ws.cell(row=row, column=9, value=_converter_decimal(diferenca))

        for col in range(5, 10):
            ws.cell(row=row, column=col).number_format = MOEDA_FORMAT

    _aplicar_borda_e_alinhamento(ws)
    _ajustar_largura_colunas(ws)
    ws.freeze_panes = "A2"


def _preencher_aba_correcoes(ws, correcoes):
    """Aba 3: Detalhamento de correções monetárias."""
    headers = [
        "Tipo",
        "Ano",
        "Mês",
        "Valor Reconhecido",
        "Valor Rec. c/ Overhead",
        "Overhead Total",
        "Diferença",
        "Taxa de Correção",
        "IGPM Acumulado (R$)",
        "Data Aplicação",
    ]

    ws.append(headers)
    _estilizar_cabecalho(ws, linha=1, cor="70AD47")

    for row, corr in enumerate(correcoes or [], start=2):
        data_correcao = corr.get("dataCorrecao")

        if data_correcao:
            try:
                dt = datetime.fromisoformat(data_correcao.replace("Z", "+00:00"))
                ano = dt.year
                mes = dt.month
            except (ValueError, TypeError, AttributeError):
                ano = ""
                mes = ""
        else:
            ano = ""
            mes = ""

        ws.cell(row=row, column=1, value=corr.get("tipo", ""))
        ws.cell(row=row, column=2, value=ano)
        ws.cell(row=row, column=3, value=mes)
        ws.cell(row=row, column=4, value=_converter_decimal(corr.get("valorReconhecido", 0)))
        ws.cell(row=row, column=5, value=_converter_decimal(corr.get("valorReconhecidoComOH", 0)))
        ws.cell(row=row, column=6, value=_converter_decimal(corr.get("overHeadTotal", 0)))
        ws.cell(row=row, column=7, value=_converter_decimal(corr.get("diferencaValor", 0)))
        ws.cell(row=row, column=8, value=_converter_decimal(corr.get("taxaCorrecao", 0)))
        ws.cell(row=row, column=9, value=_converter_decimal(corr.get("igpmAcumuladoReais", 0)))
        ws.cell(row=row, column=10, value=_formatar_data_texto(data_correcao))

        for col in [4, 5, 6, 7, 9]:
            ws.cell(row=row, column=col).number_format = MOEDA_FORMAT
        ws.cell(row=row, column=8).number_format = PERCENTUAL_FORMAT

    _aplicar_borda_e_alinhamento(ws)
    _ajustar_largura_colunas(ws)
    ws.freeze_panes = "A2"


def _preencher_aba_metadados(ws, cco):
    """Aba 4: Todos metadados da CCO."""
    ws["A1"] = "METADADOS COMPLETOS"
    ws["A1"].font = Font(size=14, bold=True, color="1F4E78")

    ws["A2"] = "Campo"
    ws["B2"] = "Valor"
    _estilizar_cabecalho(ws, linha=2, cor="A5A5A5")

    linha = 3
    for campo, valor in cco.items():
        if campo == "correcoesMonetarias":
            continue

        ws[f"A{linha}"] = campo
        ws[f"B{linha}"] = _formatar_valor_texto(valor)
        linha += 1

    _aplicar_borda_e_alinhamento(ws)
    _ajustar_largura_colunas(ws)
    ws.freeze_panes = "A3"
