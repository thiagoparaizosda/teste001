import json
from datetime import date, datetime
from decimal import Decimal

from bson import Decimal128, ObjectId
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS_CCO_PLANIFICADA = [
    # Metadados do registro
    "TIPO_REGISTRO",
    "ID_REGISTRO_PRINCIPAL",

    # Atributos principais da conta_custo_oleo_entity
    "contratoCpp",
    "campo",
    "remessa",
    "remessaExposicao",
    "faseRemessa",
    "dataReconhecimento",
    "periodo",
    "quantidadeLancamento",
    "valorLancamentoTotal",
    "valorNaoReconhecido",
    "valorReconhecido",
    "valorReconhecivel",
    "valorNaoPassivelRecuperacao",
    "valorReconhecidoExploracao",
    "valorReconhecidoProducao",
    "overHeadExploracao",
    "overHeadProducao",
    "overHeadTotal",
    "valorReconhecidoComOH",
    "origemDosGastos",
    "idRemessaGeradora",
    "flgRecuperado",
    "mesReconhecimento",
    "anoReconhecimento",
    "mesAnoReferencia",
    "faseRespostaGestora",
    "dataLancamento",
    "version",
    "_class",

    # Atributos das correcoesMonetarias
    "cm_id",
    "cm_tipo",
    "cm_subTipo",
    "cm_contrato",
    "cm_campo",
    "cm_dataCorrecao",
    "cm_dataCriacaoCorrecao",
    "cm_valorReconhecido",
    "cm_valorReconhecidoComOH",
    "cm_overHeadExploracao",
    "cm_overHeadProducao",
    "cm_overHeadTotal",
    "cm_diferencaValor",
    "cm_valorReconhecidoComOhOriginal",
    "cm_valorRecuperado",
    "cm_valorRecuperadoTotal",
    "cm_faseRemessa",
    "cm_taxaCorrecao",
    "cm_ativo",
    "cm_quantidadeLancamento",
    "cm_valorLancamentoTotal",
    "cm_valorNaoPassivelRecuperacao",
    "cm_valorReconhecivel",
    "cm_valorNaoReconhecido",
    "cm_valorReconhecidoExploracao",
    "cm_valorReconhecidoProducao",
    "cm_observacao",
    "cm_igpmAcumulado",
    "cm_igpmAcumuladoReais",
    "cm_transferencia",
    "cm_idContaCustoOleoCorrigida",
]

CAMPOS_MONETARIOS = {
    "valorLancamentoTotal",
    "valorNaoReconhecido",
    "valorReconhecido",
    "valorReconhecivel",
    "valorNaoPassivelRecuperacao",
    "valorReconhecidoExploracao",
    "valorReconhecidoProducao",
    "overHeadExploracao",
    "overHeadProducao",
    "overHeadTotal",
    "valorReconhecidoComOH",
    "cm_valorReconhecido",
    "cm_valorReconhecidoComOH",
    "cm_overHeadExploracao",
    "cm_overHeadProducao",
    "cm_overHeadTotal",
    "cm_diferencaValor",
    "cm_valorReconhecidoComOhOriginal",
    "cm_valorRecuperado",
    "cm_valorRecuperadoTotal",
    "cm_valorLancamentoTotal",
    "cm_valorNaoPassivelRecuperacao",
    "cm_valorReconhecivel",
    "cm_valorNaoReconhecido",
    "cm_valorReconhecidoExploracao",
    "cm_valorReconhecidoProducao",
    "cm_igpmAcumuladoReais",
}

CAMPOS_PERCENTUAIS = {
    "cm_taxaCorrecao",
    "cm_igpmAcumulado",
}

CAMPOS_DATA = {
    "dataReconhecimento",
    "dataLancamento",
    "cm_dataCorrecao",
    "cm_dataCriacaoCorrecao",
}

MOEDA_FORMAT = 'R$ #,##0.00'
PERCENTUAL_FORMAT = '0.0000%'
DATA_HORA_FORMAT = 'DD/MM/YYYY HH:MM:SS'


def _converter_decimal(valor, padrao=0):
    """Converte Decimal128, Decimal, int, float e strings numéricas para float."""
    if valor is None or valor == "":
        return padrao

    if isinstance(valor, Decimal128):
        return float(valor.to_decimal())

    if isinstance(valor, Decimal):
        return float(valor)

    try:
        return float(valor)
    except (TypeError, ValueError):
        return padrao


def _formatar_data_excel(valor):
    """Converte data para datetime quando possível, mantendo string inválida como texto."""
    if valor is None or valor == "":
        return ""

    if isinstance(valor, datetime):
        return valor

    if isinstance(valor, date):
        return datetime(valor.year, valor.month, valor.day)

    if isinstance(valor, str):
        texto = valor.strip()
        if not texto:
            return ""

        texto_iso = texto.replace("Z", "+00:00")
        try:
            data = datetime.fromisoformat(texto_iso)
            if data.tzinfo is not None:
                data = data.replace(tzinfo=None)
            return data
        except ValueError:
            return texto

    return str(valor)


def _valor_seguro(valor):
    """Converte valores complexos para escrita segura no Excel."""
    if valor is None:
        return ""

    if isinstance(valor, Decimal128):
        return _converter_decimal(valor)

    if isinstance(valor, Decimal):
        return float(valor)

    if isinstance(valor, ObjectId):
        return str(valor)

    if isinstance(valor, (datetime, date)):
        return _formatar_data_excel(valor)

    if isinstance(valor, (dict, list)):
        return json.dumps(valor, ensure_ascii=False, default=str)

    return valor


def _obter_campo(documento, campo, padrao=""):
    valor = documento.get(campo, padrao)
    return _valor_seguro(valor)


def _montar_linha_principal(cco):
    linha = [
        "PRINCIPAL",
        _obter_campo(cco, "_id"),
        _obter_campo(cco, "contratoCpp"),
        _obter_campo(cco, "campo"),
        _obter_campo(cco, "remessa"),
        _obter_campo(cco, "remessaExposicao"),
        _obter_campo(cco, "faseRemessa"),
        _formatar_data_excel(cco.get("dataReconhecimento")),
        _obter_campo(cco, "periodo"),
        _obter_campo(cco, "quantidadeLancamento"),
        _converter_decimal(cco.get("valorLancamentoTotal")),
        _converter_decimal(cco.get("valorNaoReconhecido")),
        _converter_decimal(cco.get("valorReconhecido")),
        _converter_decimal(cco.get("valorReconhecivel")),
        _converter_decimal(cco.get("valorNaoPassivelRecuperacao")),
        _converter_decimal(cco.get("valorReconhecidoExploracao")),
        _converter_decimal(cco.get("valorReconhecidoProducao")),
        _converter_decimal(cco.get("overHeadExploracao")),
        _converter_decimal(cco.get("overHeadProducao")),
        _converter_decimal(cco.get("overHeadTotal")),
        _converter_decimal(cco.get("valorReconhecidoComOH")),
        _obter_campo(cco, "origemDosGastos"),
        _obter_campo(cco, "idRemessaGeradora"),
        _obter_campo(cco, "flgRecuperado"),
        _obter_campo(cco, "mesReconhecimento"),
        _obter_campo(cco, "anoReconhecimento"),
        _obter_campo(cco, "mesAnoReferencia"),
        _obter_campo(cco, "faseRespostaGestora"),
        _formatar_data_excel(cco.get("dataLancamento")),
        _obter_campo(cco, "version"),
        _obter_campo(cco, "_class"),
    ]

    campos_correcoes_vazios = [""] * (len(HEADERS_CCO_PLANIFICADA) - len(linha))
    return linha + campos_correcoes_vazios


def _montar_linha_correcao(cco, correcao, indice):
    observacao = correcao.get("observacao") or correcao.get("observacoes") or ""

    linha_principal_contexto = _montar_linha_principal(cco)[:31]

    campos_correcao = [
        indice + 1,
        _obter_campo(correcao, "tipo"),
        _obter_campo(correcao, "subTipo"),
        _obter_campo(correcao, "contrato"),
        _obter_campo(correcao, "campo"),
        _formatar_data_excel(correcao.get("dataCorrecao")),
        _formatar_data_excel(correcao.get("dataCriacaoCorrecao")),
        _converter_decimal(correcao.get("valorReconhecido")),
        _converter_decimal(correcao.get("valorReconhecidoComOH")),
        _converter_decimal(correcao.get("overHeadExploracao")),
        _converter_decimal(correcao.get("overHeadProducao")),
        _converter_decimal(correcao.get("overHeadTotal")),
        _converter_decimal(correcao.get("diferencaValor")),
        _converter_decimal(correcao.get("valorReconhecidoComOhOriginal")),
        _converter_decimal(correcao.get("valorRecuperado")),
        _converter_decimal(correcao.get("valorRecuperadoTotal")),
        _obter_campo(correcao, "faseRemessa"),
        _converter_decimal(correcao.get("taxaCorrecao")),
        _obter_campo(correcao, "ativo"),
        _obter_campo(correcao, "quantidadeLancamento"),
        _converter_decimal(correcao.get("valorLancamentoTotal")),
        _converter_decimal(correcao.get("valorNaoPassivelRecuperacao")),
        _converter_decimal(correcao.get("valorReconhecivel")),
        _converter_decimal(correcao.get("valorNaoReconhecido")),
        _converter_decimal(correcao.get("valorReconhecidoExploracao")),
        _converter_decimal(correcao.get("valorReconhecidoProducao")),
        _valor_seguro(observacao),
        _converter_decimal(correcao.get("igpmAcumulado")),
        _converter_decimal(correcao.get("igpmAcumuladoReais")),
        _obter_campo(correcao, "transferencia"),
        _obter_campo(correcao, "idContaCustoOleoCorrigida"),
    ]

    linha_principal_contexto[0] = f"CORRECAO_{indice + 1}"
    return linha_principal_contexto + campos_correcao


def _aplicar_estilos(ws):
    fill_cabecalho = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fonte_cabecalho = Font(color="FFFFFF", bold=True)
    borda = Border(
        left=Side(style="thin", color="D9EAF7"),
        right=Side(style="thin", color="D9EAF7"),
        top=Side(style="thin", color="D9EAF7"),
        bottom=Side(style="thin", color="D9EAF7"),
    )

    for cell in ws[1]:
        cell.fill = fill_cabecalho
        cell.font = fonte_cabecalho
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = borda

    header_index = {header: index + 1 for index, header in enumerate(HEADERS_CCO_PLANIFICADA)}

    for row in ws.iter_rows(min_row=2):
        tipo = row[0].value
        if tipo == "PRINCIPAL":
            row[0].font = Font(bold=True, color="1F4E78")
        elif isinstance(tipo, str) and tipo.startswith("CORRECAO_"):
            row[0].font = Font(bold=True, color="70AD47")

        for cell in row:
            cell.border = borda
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for campo in CAMPOS_MONETARIOS:
        col_idx = header_index.get(campo)
        if col_idx:
            for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2):
                for item in cell:
                    item.number_format = MOEDA_FORMAT

    for campo in CAMPOS_PERCENTUAIS:
        col_idx = header_index.get(campo)
        if col_idx:
            for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2):
                for item in cell:
                    item.number_format = PERCENTUAL_FORMAT

    for campo in CAMPOS_DATA:
        col_idx = header_index.get(campo)
        if col_idx:
            for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2):
                for item in cell:
                    if isinstance(item.value, datetime):
                        item.number_format = DATA_HORA_FORMAT

    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 45)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def gerar_excel_lista_ccos_planificada(ccos):
    """
    Gera um Excel planificado com todas as CCOs retornadas pela pesquisa.

    Para cada CCO é criada:
        - uma linha PRINCIPAL;
        - uma linha CORRECAO_N para cada correção monetária existente.

    A estrutura segue o mesmo cabeçalho do script Mongo de planificação usado como referência.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "CCOs Planificadas"

    ws.append(HEADERS_CCO_PLANIFICADA)

    for cco in ccos or []:
        ws.append(_montar_linha_principal(cco))

        correcoes = cco.get("correcoesMonetarias") or []
        for indice, correcao in enumerate(correcoes):
            ws.append(_montar_linha_correcao(cco, correcao, indice))

    _aplicar_estilos(ws)

    return wb
